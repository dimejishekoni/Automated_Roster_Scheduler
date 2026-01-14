import math
from config import GridConfig
from datetime import time, datetime
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font

# ================= STRING & CLEANING UTILS =================

def clean_merge_area(ws, min_row, min_col, max_row, max_col):
    """Safely removes merges in a range to prevent overlap errors."""
    ranges_to_remove = []
    for merged_range in ws.merged_cells.ranges:
        mr_min_col, mr_min_row, mr_max_col, mr_max_row = merged_range.bounds
        if (mr_min_col <= max_col and mr_max_col >= min_col and
            mr_min_row <= max_row and mr_max_row >= min_row):
            ranges_to_remove.append(merged_range)
    for r in ranges_to_remove:
        ws.unmerge_cells(str(r))

def normalize_string(s: str) -> str:
    """Removes all whitespace for easier matching."""
    return "".join(str(s).split())

# ================= TIME & BLOCK MATH =================

def hhmm_to_minutes(hhmm: str) -> int:
    h, m = map(int, hhmm.split(":"))
    return h * 60 + m

def minutes_to_hhmm(m: int) -> str:
    m %= 24 * 60
    h = m // 60
    mm = m % 60
    return f"{h:02d}:{mm:02d}"

def to_minutes_since_start(hhmm: str, cfg) -> int:
    """
    Calculates minutes elapsed since the grid's start_hour.
    Handles shifts wrapping past midnight (e.g. 01:00 > 23:00).
    Note: 'cfg' is duck-typed (expects .start_hour) to avoid circular imports.
    """
    h, m = map(int, hhmm.split(":"))
    s = cfg.start_hour % 24
    if h < s:  # 00:xx => next day
        h += 24
    return max(0, (h - cfg.start_hour) * 60 + m)

def block_index_start(hhmm: str, cfg) -> int:
    """Calculates the starting block index (ceiling)."""
    return int(math.ceil(to_minutes_since_start(hhmm, cfg) / cfg.minutes_per_block))

def block_index_end(hhmm: str, cfg, round_finish="ceil") -> int:
    """Calculates the ending block index."""
    mins = to_minutes_since_start(hhmm, cfg)
    if round_finish == "floor":
        return int(mins // cfg.minutes_per_block)
    return int(math.ceil(mins / cfg.minutes_per_block))

def duration_minutes(start_hhmm: str, finish_hhmm: str, cfg) -> int:
    """Calculates duration in minutes between two HH:MM strings."""
    ms = to_minutes_since_start(start_hhmm, cfg)
    me = to_minutes_since_start(finish_hhmm, cfg)
    if me <= ms:
        me += 24 * 60
    return me - ms

def calculate_hourly_counts(all_staff_lists, cfg: GridConfig):
    """
    Returns a list of integer counts (one per hour) representing 
    how many staff are active during that hour across all sections.
    """
    counts = [0] * cfg.total_hours
    all_staff = [s for sublist in all_staff_lists for s in sublist]
    
    for s in all_staff:
        start_min = to_minutes_since_start(s["start_time"], cfg)
        end_min = to_minutes_since_start(s["finish_time"], cfg)
        
        if end_min <= start_min: 
            end_min += 24 * 60

        for h in range(cfg.total_hours):
            hour_start = h * 60
            hour_end = (h + 1) * 60
            if max(start_min, hour_start) < min(end_min, hour_end):
                counts[h] += 1
    return counts

# ================= NEW Helper: Draw Bold Outline =================
def draw_section_outline(ws, row_start, row_end, cfg, top_style="thick"):
    """
    Draws a THICK BLACK border around the rectangular grid of the section.
    (Top of row_start, Bottom of row_end, Left of start_col, Right of end_col)
    """
    thick = Side(border_style="thick", color="000000")
    thin  = Side(border_style="thin",  color="000000")
    top_side = thick if top_style == "thick" else thin

    col_start = 1 
    col_end = cfg.start_column + cfg.total_blocks - 1

    # 1. Top Edge
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row_start, column=c)
        current = cell.border
        cell.border = Border(top=top_side, bottom=current.bottom, left=current.left, right=current.right)

    # 2. Bottom Edge (Always Thick)
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row_end, column=c)
        current = cell.border
        cell.border = Border(top=current.top, bottom=thick, left=current.left, right=current.right)

    # 3. Left Edge (Always Thick)
    for r in range(row_start, row_end + 1):
        cell = ws.cell(row=r, column=col_start)
        current = cell.border
        cell.border = Border(top=current.top, bottom=current.bottom, left=thick, right=current.right)

    # 4. Right Edge (Always Thick)
    for r in range(row_start, row_end + 1):
        cell = ws.cell(row=r, column=col_end)
        current = cell.border
        cell.border = Border(top=current.top, bottom=current.bottom, left=current.left, right=thick)
        
    # 5. Fix Corners
    # Top-Left
    tl = ws.cell(row=row_start, column=col_start)
    tl.border = Border(top=top_side, left=thick, bottom=tl.border.bottom, right=tl.border.right)
    # Top-Right
    tr = ws.cell(row=row_start, column=col_end)
    tr.border = Border(top=top_side, right=thick, bottom=tr.border.bottom, left=tr.border.left)
    # Bottom-Left
    bl = ws.cell(row=row_end, column=col_start)
    bl.border = Border(bottom=thick, left=thick, top=bl.border.top, right=bl.border.right)
    # Bottom-Right
    br = ws.cell(row=row_end, column=col_end)
    br.border = Border(bottom=thick, right=thick, top=br.border.top, left=br.border.left)
